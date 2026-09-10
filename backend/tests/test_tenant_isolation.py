"""Tenant isolation — every tenant-scoped router must neither leak nor touch another tenant's rows.

Conventions used here:
- "A" is the caller, "B" is the victim. A must never see or change B's data.
- Objects addressed by id that belong to B must produce 404 (not 403 — that would confirm existence).
- tenant_id is always derived from the JWT; any tenant_id in a request body is ignored.
"""

from __future__ import annotations

import io
import json
import zipfile
from dataclasses import dataclass

import pytest
from sqlalchemy import select

from app.models.booking import Booking
from app.models.correction import Correction
from app.models.kontenplan import Konto
from app.models.memory import Memory
from app.models.review_queue import ReviewQueueItem
from app.models.scanner_config import ScannerConfig
from app.models.tenant import Tenant
from app.models.training_data import TrainingRow
from app.models.user import User
from app.services.classifier import make_memory_key
from app.services.receipts import store_receipt
from app.services.scanner.scanner_service import ScannerService
from tests.factories import (
    auth_headers,
    create_audit_entry,
    create_booking,
    create_correction,
    create_konto,
    create_konto_default,
    create_memory,
    create_review_item,
    create_scanner_config,
    create_tenant,
    create_training_row,
    create_user,
)


@dataclass
class TwoTenants:
    tenant_a: Tenant
    user_a: User
    headers_a: dict[str, str]
    tenant_b: Tenant
    user_b: User
    headers_b: dict[str, str]


@pytest.fixture
async def tenants(db_session) -> TwoTenants:
    tenant_a = await create_tenant(db_session, "A")
    tenant_b = await create_tenant(db_session, "B")
    user_a = await create_user(db_session, tenant_a)
    user_b = await create_user(db_session, tenant_b)
    return TwoTenants(tenant_a, user_a, auth_headers(user_a), tenant_b, user_b, auth_headers(user_b))


async def _rows(db_session, model, tenant_id: int):
    result = await db_session.execute(select(model).where(model.tenant_id == tenant_id))
    return result.scalars().all()


# ── service-level (pre-existing) ─────────────────────────────────────────────


@pytest.mark.asyncio
async def test_scanner_config_is_tenant_scoped(db_session):
    tenant_a = await create_tenant(db_session, "A")
    tenant_b = await create_tenant(db_session, "B")
    user_a = await create_user(db_session, tenant_a)
    user_b = await create_user(db_session, tenant_b)

    config_a = await ScannerService(db_session, user_a).get_or_create_config_model()
    config_b = await ScannerService(db_session, user_b).get_or_create_config_model()

    assert config_a.tenant_id == tenant_a.id
    assert config_b.tenant_id == tenant_b.id
    assert config_a.id != config_b.id


@pytest.mark.asyncio
async def test_config_query_never_returns_other_tenant(db_session):
    tenant_a = await create_tenant(db_session, "A")
    tenant_b = await create_tenant(db_session, "B")
    user_b = await create_user(db_session, tenant_b)

    await ScannerService(db_session, user_b).get_or_create_config_model()

    stmt = select(ScannerConfig).where(ScannerConfig.tenant_id == tenant_a.id)
    result = await db_session.execute(stmt)
    assert result.scalar_one_or_none() is None


@pytest.mark.asyncio
async def test_bookings_are_isolated_by_tenant(db_session):
    tenant_a = await create_tenant(db_session, "A")
    tenant_b = await create_tenant(db_session, "B")

    await create_booking(db_session, tenant_a, beschreibung="Tenant A invoice")

    stmt = select(Booking).where(Booking.tenant_id == tenant_b.id)
    result = await db_session.execute(stmt)
    assert result.scalars().all() == []


# ── auth guard ───────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("method", "path"),
    [
        ("GET", "/api/bookings/"),
        ("GET", "/api/bookings/stats"),
        ("GET", "/api/bookings/1/source"),
        ("GET", "/api/review/"),
        ("GET", "/api/kontenplan/"),
        ("GET", "/api/kontenplan/defaults"),
        ("GET", "/api/export/banana"),
        ("GET", "/api/export/csv"),
        ("GET", "/api/export/excel"),
        ("POST", "/api/export/banana"),
        ("POST", "/api/export/csv"),
        ("POST", "/api/export/excel"),
        ("POST", "/api/export/email/rows"),
        ("GET", "/api/stats/learning"),
        ("GET", "/api/audit/"),
        ("GET", "/api/scanner/config"),
        ("GET", "/api/classify/info"),
        ("GET", "/api/classify/memory"),
        ("GET", "/api/classify/corrections"),
        ("GET", "/api/classify/top-classes"),
        ("GET", "/api/classify/download/memory"),
        ("DELETE", "/api/classify/memory"),
        ("POST", "/api/classify/train"),
    ],
)
@pytest.mark.asyncio
async def test_tenant_endpoints_reject_anonymous(client, method, path):
    resp = await client.request(method, path)
    assert resp.status_code in (401, 403)
    assert "error" in resp.json()


# ── bookings ─────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_bookings_list_and_stats_only_own(client, db_session, tenants: TwoTenants):
    await create_booking(db_session, tenants.tenant_a, beschreibung="A-1", betrag=10.0)
    await create_booking(db_session, tenants.tenant_b, beschreibung="B-1", betrag=999.0)
    await create_booking(db_session, tenants.tenant_b, beschreibung="B-2", betrag=1.0)

    resp = await client.get("/api/bookings/", headers=tenants.headers_a)
    assert resp.status_code == 200
    assert [b["beschreibung"] for b in resp.json()] == ["A-1"]

    resp = await client.get("/api/bookings/stats", headers=tenants.headers_a)
    assert resp.status_code == 200
    assert resp.json()["total_count"] == 1
    assert resp.json()["total_amount"] == 10.0


@pytest.mark.asyncio
async def test_booking_create_ignores_tenant_id_in_body(client, db_session, tenants: TwoTenants):
    payload = [{"beschreibung": "smuggled", "betrag": 5.0, "tenant_id": tenants.tenant_b.id}]
    resp = await client.post("/api/bookings/", json=payload, headers=tenants.headers_a)
    assert resp.status_code == 200
    await db_session.commit()

    assert [b.beschreibung for b in await _rows(db_session, Booking, tenants.tenant_a.id)] == ["smuggled"]
    assert await _rows(db_session, Booking, tenants.tenant_b.id) == []


@pytest.mark.asyncio
async def test_booking_source_document_is_404_for_other_tenant(client, db_session, tenants: TwoTenants):
    key_b = store_receipt(tenants.tenant_b.id, filename="b.pdf", content_type="application/pdf", content=b"%PDF-B")
    booking_b = await create_booking(db_session, tenants.tenant_b, beschreibung="B-Beleg")
    booking_b.source_key = key_b
    await db_session.commit()

    # By id: B's booking does not exist for A.
    resp = await client.get(f"/api/bookings/{booking_b.id}/source", headers=tenants.headers_a)
    assert resp.status_code == 404
    assert resp.json()["error"]["code"]

    # By key: A cannot attach B's document to its own booking and read it that way.
    resp = await client.post(
        "/api/bookings/", json=[{"beschreibung": "steal", "source_key": key_b}], headers=tenants.headers_a
    )
    assert resp.status_code == 400
    assert await _rows(db_session, Booking, tenants.tenant_a.id) == []

    # B itself still gets the file.
    resp = await client.get(f"/api/bookings/{booking_b.id}/source", headers=tenants.headers_b)
    assert resp.status_code == 200
    assert resp.content == b"%PDF-B"


# ── review queue ─────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_review_list_only_own(client, db_session, tenants: TwoTenants):
    await create_review_item(db_session, tenants.tenant_a, beschreibung="A pending")
    await create_review_item(db_session, tenants.tenant_b, beschreibung="B pending")

    resp = await client.get("/api/review/", headers=tenants.headers_a)
    assert resp.status_code == 200
    assert [i["beschreibung"] for i in resp.json()["items"]] == ["A pending"]
    assert resp.json()["count"] == 1


@pytest.mark.parametrize("action", ["approve", "reject"])
@pytest.mark.asyncio
async def test_review_foreign_item_is_404_and_untouched(client, db_session, tenants: TwoTenants, action):
    item_b = await create_review_item(db_session, tenants.tenant_b)

    resp = await client.post(f"/api/review/{item_b.id}/{action}", json={}, headers=tenants.headers_a)
    assert resp.status_code == 404

    await db_session.refresh(item_b)
    assert item_b.status == "pending"
    assert item_b.resolved_at is None
    # And no memory got written for the wrong tenant either.
    assert await _rows(db_session, Memory, tenants.tenant_a.id) == []
    assert await _rows(db_session, Memory, tenants.tenant_b.id) == []


@pytest.mark.asyncio
async def test_review_approve_own_item_writes_only_own_memory(client, db_session, tenants: TwoTenants):
    item_a = await create_review_item(db_session, tenants.tenant_a, beschreibung="Coop Einkauf")

    resp = await client.post(
        f"/api/review/{item_a.id}/approve", json={"corrected_soll": "6510"}, headers=tenants.headers_a
    )
    assert resp.status_code == 200

    await db_session.refresh(item_a)
    assert item_a.status == "approved"
    assert item_a.resolved_soll == "6510"
    mem_a = await _rows(db_session, Memory, tenants.tenant_a.id)
    assert [m.kt_soll for m in mem_a] == ["6510"]
    assert await _rows(db_session, Memory, tenants.tenant_b.id) == []


# ── scanner config ───────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_scanner_config_get_returns_own_tenant(client, db_session, tenants: TwoTenants):
    await create_scanner_config(db_session, tenants.tenant_b, ollama_base_url="http://b-only:11434")

    resp = await client.get("/api/scanner/config", headers=tenants.headers_a)
    assert resp.status_code == 200
    assert resp.json()["tenant_id"] == tenants.tenant_a.id
    assert resp.json()["ollama_base_url"] != "http://b-only:11434"


@pytest.mark.asyncio
async def test_scanner_config_patch_does_not_touch_other_tenant(client, db_session, tenants: TwoTenants):
    config_b = await create_scanner_config(db_session, tenants.tenant_b, review_confidence_threshold=0.9)

    resp = await client.patch(
        "/api/scanner/config", json={"review_confidence_threshold": 0.1}, headers=tenants.headers_a
    )
    assert resp.status_code == 200
    assert resp.json()["review_confidence_threshold"] == 0.1

    await db_session.refresh(config_b)
    assert config_b.review_confidence_threshold == 0.9
    configs_a = await _rows(db_session, ScannerConfig, tenants.tenant_a.id)
    assert len(configs_a) == 1 and configs_a[0].review_confidence_threshold == 0.1


@pytest.mark.asyncio
async def test_scanner_config_put_does_not_touch_other_tenant(client, db_session, tenants: TwoTenants):
    config_b = await create_scanner_config(db_session, tenants.tenant_b, ocr_provider="b-ocr")

    payload = {
        "ocr_provider": "a-ocr",
        "vision_provider": "ollama",
        "ollama_base_url": "http://a:11434",
    }
    resp = await client.put("/api/scanner/config", json=payload, headers=tenants.headers_a)
    assert resp.status_code == 200
    assert resp.json()["tenant_id"] == tenants.tenant_a.id
    assert resp.json()["ocr_provider"] == "a-ocr"

    await db_session.refresh(config_b)
    assert config_b.ocr_provider == "b-ocr"


# ── export ───────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("fmt", ["banana", "csv", "excel"])
@pytest.mark.asyncio
async def test_export_only_own_bookings(client, db_session, tenants: TwoTenants, fmt):
    await create_booking(db_session, tenants.tenant_a, beschreibung="ONLY-A-ROW")
    await create_booking(db_session, tenants.tenant_b, beschreibung="ONLY-B-ROW")

    resp = await client.get(f"/api/export/{fmt}", headers=tenants.headers_a)
    assert resp.status_code == 200
    body = resp.content
    if fmt == "excel":
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(body)).active
        cells = {str(c.value) for row in ws.iter_rows() for c in row if c.value is not None}
        assert "ONLY-A-ROW" in cells
        assert "ONLY-B-ROW" not in cells
    else:
        assert b"ONLY-A-ROW" in body
        assert b"ONLY-B-ROW" not in body


@pytest.mark.parametrize("fmt", ["banana", "csv", "excel"])
@pytest.mark.asyncio
async def test_export_empty_for_tenant_without_bookings(client, db_session, tenants: TwoTenants, fmt):
    await create_booking(db_session, tenants.tenant_b, beschreibung="ONLY-B-ROW")

    resp = await client.get(f"/api/export/{fmt}", headers=tenants.headers_a)
    assert resp.status_code == 404
    assert resp.json()["error"]["code"] == "http_404"


@pytest.mark.skip(reason="POST /api/export/email needs a configured SMTP server; the DB read path is covered above.")
async def test_export_email_only_own_bookings():
    pass


# ── kontenplan ───────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_kontenplan_get_only_own(client, db_session, tenants: TwoTenants):
    await create_konto(db_session, tenants.tenant_a, "6500", "A Büro")
    await create_konto(db_session, tenants.tenant_b, "6500", "B Büro")
    await create_konto(db_session, tenants.tenant_b, "4000", "B Material")

    resp = await client.get("/api/kontenplan/", headers=tenants.headers_a)
    assert resp.status_code == 200
    assert resp.json() == {"kontenplan": {"6500": "A Büro"}}


@pytest.mark.asyncio
async def test_kontenplan_put_replaces_only_own(client, db_session, tenants: TwoTenants):
    await create_konto(db_session, tenants.tenant_a, "6500", "A old")
    await create_konto(db_session, tenants.tenant_b, "6500", "B keep")

    resp = await client.put(
        "/api/kontenplan/", json={"kontenplan": {"1020": "Bank", "6570": "IT"}}, headers=tenants.headers_a
    )
    assert resp.status_code == 200
    assert resp.json()["count"] == 2

    rows_a = {k.konto_nr: k.beschreibung for k in await _rows(db_session, Konto, tenants.tenant_a.id)}
    rows_b = {k.konto_nr: k.beschreibung for k in await _rows(db_session, Konto, tenants.tenant_b.id)}
    assert rows_a == {"1020": "Bank", "6570": "IT"}
    assert rows_b == {"6500": "B keep"}


@pytest.mark.asyncio
async def test_kontenplan_defaults_only_own(client, db_session, tenants: TwoTenants):
    await create_konto_default(db_session, tenants.tenant_a, "6500", "1020", "I81", "8.10")
    await create_konto_default(db_session, tenants.tenant_b, "4000", "1000", "M81", "8.10")

    resp = await client.get("/api/kontenplan/defaults", headers=tenants.headers_a)
    assert resp.status_code == 200
    assert list(resp.json()["defaults"]) == ["6500"]


# ── stats ────────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_learning_stats_only_own(client, db_session, tenants: TwoTenants):
    await create_memory(db_session, tenants.tenant_a, "Coop", kt_soll="6500")
    await create_memory(db_session, tenants.tenant_b, "Migros", kt_soll="6500")
    await create_memory(db_session, tenants.tenant_b, "Swisscom", kt_soll="6510")
    await create_correction(db_session, tenants.tenant_b)
    await create_booking(db_session, tenants.tenant_b, source="scanner")

    resp = await client.get("/api/stats/learning", headers=tenants.headers_a)
    assert resp.status_code == 200
    data = resp.json()
    assert data["memory_count"] == 1
    assert data["correction_count"] == 0
    assert data["booking_count"] == 0
    assert data["memory_distribution"] == [{"account": "6500", "count": 1}]
    assert data["correction_distribution"] == []
    assert data["source_distribution"] == []


# ── audit ────────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_audit_list_only_own(client, db_session, tenants: TwoTenants):
    await create_audit_entry(db_session, tenants.tenant_a, action="a.only")
    await create_audit_entry(db_session, tenants.tenant_b, action="b.only", actor_user_id=tenants.user_b.id)

    resp = await client.get("/api/audit/", headers=tenants.headers_a)
    assert resp.status_code == 200
    assert [e["action"] for e in resp.json()["items"]] == ["a.only"]


# ── import ───────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_import_banana_csv_writes_only_own_tenant(client, db_session, tenants: TwoTenants):
    await create_training_row(db_session, tenants.tenant_b, "B existing", "6500")
    csv = "Beschreibung;KtSoll;KtHaben;MwStCode;MwStPct\nCoop Einkauf;6500;1020;;\nSwisscom Abo;6510;1020;I81;8.10\n"

    resp = await client.post(
        "/api/import/banana",
        params={"replace": "true", "also_memory": "true", "auto_train": "false"},
        files={"file": ("export.csv", csv.encode(), "text/csv")},
        headers=tenants.headers_a,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["imported"] == 2

    rows_a = await _rows(db_session, TrainingRow, tenants.tenant_a.id)
    rows_b = await _rows(db_session, TrainingRow, tenants.tenant_b.id)
    assert sorted(r.beschreibung for r in rows_a) == ["Coop Einkauf", "Swisscom Abo"]
    # replace=true must only wipe the caller's rows.
    assert [r.beschreibung for r in rows_b] == ["B existing"]

    mem_a = await _rows(db_session, Memory, tenants.tenant_a.id)
    assert sorted(m.lookup_key for m in mem_a) == sorted(
        [make_memory_key("Coop Einkauf"), make_memory_key("Swisscom Abo")]
    )
    assert await _rows(db_session, Memory, tenants.tenant_b.id) == []


# ── classify ─────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("path", ["/api/classify/", "/api/classify/predict"])
@pytest.mark.asyncio
async def test_classify_never_uses_other_tenants_memory(client, db_session, tenants: TwoTenants, path):
    text = "Zauberladen Rechnung 4711"
    await create_memory(db_session, tenants.tenant_b, text, kt_soll="1234", kt_haben="4321")

    resp = await client.post(path, json={"beschreibung": text, "betrag": 100.0}, headers=tenants.headers_a)
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["source"] != "Gedächtnis"
    assert data["kt_soll"] != "1234"

    resp = await client.post(path, json={"beschreibung": text, "betrag": 100.0}, headers=tenants.headers_b)
    assert resp.json()["source"] == "Gedächtnis"
    assert resp.json()["kt_soll"] == "1234"


@pytest.mark.asyncio
async def test_classify_low_confidence_enqueues_review_for_caller_only(client, db_session, tenants: TwoTenants):
    resp = await client.post(
        "/api/classify/", json={"beschreibung": "xqzv unbekannt", "betrag": 12.0}, headers=tenants.headers_a
    )
    assert resp.status_code == 200
    assert resp.json()["needs_review"] is True

    items_a = await _rows(db_session, ReviewQueueItem, tenants.tenant_a.id)
    assert len(items_a) == 1 and items_a[0].id == resp.json()["review_id"]
    assert await _rows(db_session, ReviewQueueItem, tenants.tenant_b.id) == []


@pytest.mark.asyncio
async def test_classify_batch_never_uses_other_tenants_memory(client, db_session, tenants: TwoTenants):
    await create_memory(db_session, tenants.tenant_b, "Fremdfirma", kt_soll="1234")

    resp = await client.post(
        "/api/classify/batch",
        json={"transactions": [{"Beschreibung": "Fremdfirma", "Betrag CHF": 50}]},
        headers=tenants.headers_a,
    )
    assert resp.status_code == 200
    assert resp.json()["results"][0]["kt_soll"] != "1234"


@pytest.mark.parametrize(
    ("action", "model"),
    [("memory", Memory), ("corrections", Correction)],
)
@pytest.mark.asyncio
async def test_classify_delete_only_own(client, db_session, tenants: TwoTenants, action, model):
    if model is Memory:
        await create_memory(db_session, tenants.tenant_a, "A")
        await create_memory(db_session, tenants.tenant_b, "B")
    else:
        await create_correction(db_session, tenants.tenant_a)
        await create_correction(db_session, tenants.tenant_b)

    resp = await client.delete(f"/api/classify/{action}", headers=tenants.headers_a)
    assert resp.status_code == 200

    assert await _rows(db_session, model, tenants.tenant_a.id) == []
    assert len(await _rows(db_session, model, tenants.tenant_b.id)) == 1


@pytest.mark.asyncio
async def test_classify_lists_and_info_only_own(client, db_session, tenants: TwoTenants):
    await create_memory(db_session, tenants.tenant_a, "A shop")
    await create_memory(db_session, tenants.tenant_b, "B shop")
    await create_correction(db_session, tenants.tenant_b, beschreibung="B corr")
    await create_training_row(db_session, tenants.tenant_b, "B train", "4000")

    resp = await client.get("/api/classify/memory", headers=tenants.headers_a)
    assert [m["lookup_key"] for m in resp.json()["entries"]] == [make_memory_key("A shop")]

    resp = await client.get("/api/classify/corrections", headers=tenants.headers_a)
    assert resp.json() == {"corrections": [], "count": 0}

    resp = await client.get("/api/classify/top-classes", headers=tenants.headers_a)
    assert resp.json() == []

    resp = await client.get("/api/classify/info", headers=tenants.headers_a)
    assert resp.json()["memory_count"] == 1
    assert resp.json()["correction_count"] == 0
    assert resp.json()["has_model"] is False


@pytest.mark.asyncio
async def test_classify_download_only_own(client, db_session, tenants: TwoTenants):
    await create_memory(db_session, tenants.tenant_a, "A shop")
    await create_memory(db_session, tenants.tenant_b, "B shop")
    await create_correction(db_session, tenants.tenant_b, beschreibung="B corr")
    await create_konto_default(db_session, tenants.tenant_b, "4000")

    resp = await client.get("/api/classify/download/memory", headers=tenants.headers_a)
    assert resp.status_code == 200
    assert [m["lookup_key"] for m in resp.json()] == [make_memory_key("A shop")]

    resp = await client.get("/api/classify/download/model", headers=tenants.headers_a)
    assert resp.status_code == 404

    resp = await client.get("/api/classify/download/bundle", headers=tenants.headers_a)
    assert resp.status_code == 200
    with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
        memory = json.loads(zf.read("memory.json"))
        corrections = json.loads(zf.read("corrections.json"))
        kontenplan = json.loads(zf.read("kontenplan.json"))
        assert "model.pkl" not in zf.namelist()
    assert [m["lookup_key"] for m in memory] == [make_memory_key("A shop")]
    assert corrections == []
    assert kontenplan == []


@pytest.mark.asyncio
async def test_classify_upload_memory_replaces_only_own(client, db_session, tenants: TwoTenants):
    await create_memory(db_session, tenants.tenant_a, "A old")
    await create_memory(db_session, tenants.tenant_b, "B keep")
    payload = json.dumps([{"lookup_key": "a new", "kt_soll": "6570", "kt_haben": "1020"}]).encode()

    resp = await client.post(
        "/api/classify/upload",
        files={"file": ("memory.json", payload, "application/json")},
        headers=tenants.headers_a,
    )
    assert resp.status_code == 200
    assert resp.json()["restored"] == ["memory"]

    assert [m.lookup_key for m in await _rows(db_session, Memory, tenants.tenant_a.id)] == ["a new"]
    assert [m.lookup_key for m in await _rows(db_session, Memory, tenants.tenant_b.id)] == [make_memory_key("B keep")]


@pytest.mark.asyncio
async def test_classify_correct_writes_only_own(client, db_session, tenants: TwoTenants):
    body = {
        "beschreibung": "Galaxus Bestellung",
        "original_soll": "6500",
        "original_haben": "1020",
        "corrected_soll": "6570",
        "corrected_haben": "1020",
    }
    resp = await client.post("/api/classify/correct", json=body, headers=tenants.headers_a)
    assert resp.status_code == 200

    assert len(await _rows(db_session, Correction, tenants.tenant_a.id)) == 1
    assert len(await _rows(db_session, Memory, tenants.tenant_a.id)) == 1
    assert await _rows(db_session, Correction, tenants.tenant_b.id) == []
    assert await _rows(db_session, Memory, tenants.tenant_b.id) == []


@pytest.mark.asyncio
async def test_classify_train_never_sees_other_tenants_data(client, db_session, tenants: TwoTenants):
    for i in range(6):
        await create_training_row(db_session, tenants.tenant_b, f"B row {i}", "6500" if i % 2 else "4000")

    resp = await client.post("/api/classify/train", headers=tenants.headers_a)
    assert resp.status_code == 400
    assert "Zu wenige Daten" in resp.json()["error"]["message"]


@pytest.mark.skip(reason="/api/scanner/extract, /api/scanner/status and /api/ai/* need a live Ollama/Tesseract.")
async def test_scanner_and_ai_endpoints_only_own():
    pass
