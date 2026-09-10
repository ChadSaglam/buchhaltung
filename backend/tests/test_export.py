"""Export — Banana TSV, semicolon CSV, styled Excel, and the HTTP layer around them.

Banana's import expects: one header line, tab-separated columns
(Date, Description, AccountDebit, AccountCredit, Amount, VatCode), ISO dates,
plain amounts with exactly two decimals and no thousands separator.
"""

from __future__ import annotations

import io
from decimal import Decimal

import pandas as pd
import pytest
from openpyxl import load_workbook

from app.services.export import df_to_banana_tsv, df_to_csv, df_to_styled_excel, fmt_swiss, round_chf
from tests.factories import auth_headers, create_booking, create_tenant, create_user

BANANA_HEADER = "Date\tDescription\tAccountDebit\tAccountCredit\tAmount\tVatCode"

COLUMNS = [
    "Nr",
    "Datum",
    "Beleg",
    "Rechnung",
    "Beschreibung",
    "KtSoll",
    "KtHaben",
    "Betrag CHF",
    "MwStUSt-Code",
    "Art Betrag",
    "MwSt-%",
    "Gebuchte MwStUSt CHF",
    "KS3",
]


def _df(*rows: dict) -> pd.DataFrame:
    base = dict.fromkeys(COLUMNS, "")
    return pd.DataFrame([{**base, "Nr": i + 1, **r} for i, r in enumerate(rows)], columns=COLUMNS)


def _row(**overrides) -> dict:
    row = {
        "Datum": "15.03.2025",
        "Beschreibung": "Büromaterial",
        "KtSoll": "6500",
        "KtHaben": "1020",
        "Betrag CHF": 12.5,
        "MwStUSt-Code": "I81",
        "MwSt-%": "8.10",
        "Gebuchte MwStUSt CHF": 0.94,
    }
    row.update(overrides)
    return row


# ── fmt_swiss ────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (0, "0.00"),
        (12.5, "12.50"),
        (1234.56, "1’234.56"),
        (1234567.89, "1’234’567.89"),
        (-1234.56, "-1’234.56"),
        (-0.5, "-0.50"),
        ("99.9", "99.90"),
        (100, "100.00"),
    ],
)
def test_fmt_swiss_formats_with_apostrophe_thousands(value, expected):
    assert fmt_swiss(value) == expected


@pytest.mark.parametrize("value", [None, "", float("nan")])
def test_fmt_swiss_blank_for_missing(value):
    assert fmt_swiss(value) == ""


# Regression guards for the money-rounding bug: the fraction used to be
# rounded on its own, so 1234.999 rendered as "1'234.00" (the carry was lost),
# and exact halves rounded half-even ("0.125" -> "0.12") instead of half-up.
@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (1234.999, "1’235.00"),  # carry into the integer part
        (999.995, "1’000.00"),  # carry across a thousands boundary
        (1_000_000.995, "1’000’001.00"),  # large amount, carry
        (0.005, "0.01"),  # half-up on a 3-decimal input
        (0.125, "0.13"),  # exact binary half: must not be half-even
        (2.675, "2.68"),  # classic float trap
        (-0.125, "-0.13"),  # symmetric for negatives
        (-0.004, "0.00"),  # no "-0.00"
    ],
)
def test_fmt_swiss_rounds_half_up_with_carry(value, expected):
    assert fmt_swiss(value) == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (Decimal("2.675"), Decimal("2.68")),
        (0.125, Decimal("0.13")),
        (-0.125, Decimal("-0.13")),
        (1234.999, Decimal("1235.00")),
        (0, Decimal("0.00")),
        ("7.777", Decimal("7.78")),
        (0.135, Decimal("0.14")),
        (-0.135, Decimal("-0.14")),
        (-2.675, Decimal("-2.68")),
        (-1234.999, Decimal("-1235.00")),
        (-0.004, Decimal("0.00")),  # no negative zero
        (100 * 8.1 / 108.1, Decimal("7.49")),  # 8.1% tax portion of CHF 100
        (100 * 2.6 / 102.6, Decimal("2.53")),  # 2.6% tax portion of CHF 100
        (Decimal("0.005"), Decimal("0.01")),
    ],
)
def test_round_chf_half_up(value, expected):
    assert round_chf(value) == expected


# ── Banana TSV ───────────────────────────────────────────────────────────────


def test_banana_tsv_header_and_tab_separated_row():
    out = df_to_banana_tsv(_df(_row()))
    lines = out.split("\n")
    assert lines[0] == BANANA_HEADER
    assert lines[1] == "2025-03-15\tBüromaterial\t6500\t1020\t12.50\tI81"
    assert len(lines) == 2
    assert not out.endswith("\n")


def test_banana_tsv_converts_swiss_date_to_iso():
    out = df_to_banana_tsv(_df(_row(Datum="01.12.2024")))
    assert out.split("\n")[1].startswith("2024-12-01\t")


def test_banana_tsv_leaves_iso_and_blank_dates_alone():
    out = df_to_banana_tsv(_df(_row(Datum="2024-12-01"), _row(Datum="")))
    rows = out.split("\n")[1:]
    assert rows[0].startswith("2024-12-01\t")
    assert rows[1].startswith("\t")


def test_banana_tsv_negative_amount_keeps_sign():
    out = df_to_banana_tsv(_df(_row(**{"Betrag CHF": -87.3})))
    assert out.split("\n")[1].split("\t")[4] == "-87.30"


def test_banana_tsv_three_decimal_input_is_rounded_to_two():
    out = df_to_banana_tsv(_df(_row(**{"Betrag CHF": 10.994})))
    assert out.split("\n")[1].split("\t")[4] == "10.99"


@pytest.mark.parametrize(
    ("value", "expected"),
    [(0.125, "0.13"), (2.675, "2.68"), (-0.125, "-0.13"), (0.005, "0.01"), (1234.999, "1235.00")],
)
def test_banana_tsv_amount_rounds_half_up(value, expected):
    # Regression guard: amounts used to go through f"{x:.2f}" (half-even on exact halves).
    out = df_to_banana_tsv(_df(_row(**{"Betrag CHF": value})))
    assert out.split("\n")[1].split("\t")[4] == expected


def test_banana_tsv_large_amount_has_no_thousands_separator():
    out = df_to_banana_tsv(_df(_row(**{"Betrag CHF": 1234567.891})))
    assert out.split("\n")[1].split("\t")[4] == "1234567.89"


def test_banana_tsv_zero_amount_is_blank():
    # Documented current behaviour: 0 is written as an empty Amount cell.
    out = df_to_banana_tsv(_df(_row(**{"Betrag CHF": 0})))
    assert out.split("\n")[1].split("\t")[4] == ""


def test_banana_tsv_unparseable_amount_is_blank():
    out = df_to_banana_tsv(_df(_row(**{"Betrag CHF": "abc"})))
    assert out.split("\n")[1].split("\t")[4] == ""


@pytest.mark.parametrize("code", ["I81", "V81", "M81", "I25", ""])
def test_banana_tsv_passes_vat_code_through(code):
    out = df_to_banana_tsv(_df(_row(**{"MwStUSt-Code": code})))
    assert out.split("\n")[1].split("\t")[5] == code


def test_banana_tsv_row_order_preserved():
    out = df_to_banana_tsv(_df(_row(Beschreibung="erste"), _row(Beschreibung="zweite"), _row(Beschreibung="dritte")))
    descs = [line.split("\t")[1] for line in out.split("\n")[1:]]
    assert descs == ["erste", "zweite", "dritte"]


def test_banana_tsv_empty_frame_is_header_only():
    assert df_to_banana_tsv(pd.DataFrame([])) == BANANA_HEADER
    assert df_to_banana_tsv(pd.DataFrame(columns=COLUMNS)) == BANANA_HEADER


# ── CSV ──────────────────────────────────────────────────────────────────────


def test_csv_is_semicolon_separated_without_index():
    out = df_to_csv(_df(_row(), _row(Beschreibung="zwei")))
    lines = out.strip().split("\n")
    assert lines[0] == ";".join(COLUMNS)
    assert len(lines) == 3
    assert lines[1].startswith("1;15.03.2025;;;Büromaterial;6500;1020;12.5;I81;")
    assert lines[2].split(";")[0] == "2"


def test_csv_empty_frame_has_only_header():
    out = df_to_csv(pd.DataFrame(columns=COLUMNS))
    assert out.strip() == ";".join(COLUMNS)


# ── Excel ────────────────────────────────────────────────────────────────────


def test_excel_writes_numbers_as_numbers_and_flags_negatives():
    data = df_to_styled_excel(_df(_row(**{"Betrag CHF": 1234.5}), _row(**{"Betrag CHF": -20})))
    ws = load_workbook(io.BytesIO(data)).active

    assert [c.value for c in ws[1]] == COLUMNS
    betrag_col = COLUMNS.index("Betrag CHF") + 1
    pos = ws.cell(row=2, column=betrag_col)
    neg = ws.cell(row=3, column=betrag_col)
    assert pos.value == 1234.5 and isinstance(pos.value, float)
    assert pos.number_format == "#,##0.00"
    assert neg.value == -20.0
    assert neg.font.color.rgb.endswith("FF0000")
    assert not pos.font.color or not (pos.font.color.rgb or "").endswith("FF0000")
    assert ws.freeze_panes == "A2"


def test_excel_blank_strings_become_empty_cells():
    data = df_to_styled_excel(_df(_row(Beleg="")))
    ws = load_workbook(io.BytesIO(data)).active
    assert ws.cell(row=2, column=COLUMNS.index("Beleg") + 1).value is None


# ── HTTP layer ───────────────────────────────────────────────────────────────


def _api_row(**overrides) -> dict:
    row = {
        "nr": 1,
        "datum": "15.03.2025",
        "beleg": "",
        "rechnung": "",
        "beschreibung": "Café Müller",
        "kt_soll": "6500",
        "kt_haben": "1020",
        "betrag": 12.5,
        "mwstcode": "I81",
        "artbetrag": "",
        "mwstpct": "8.10",
        "mwstchf": 0.94,
        "ks3": "",
    }
    row.update(overrides)
    return row


@pytest.fixture
async def headers(db_session) -> dict[str, str]:
    """The stateless POST exports read no tenant data but still require a login (B-06)."""
    tenant = await create_tenant(db_session)
    user = await create_user(db_session, tenant)
    return auth_headers(user)


@pytest.mark.parametrize(
    "path", ["/api/export/banana", "/api/export/csv", "/api/export/excel", "/api/export/email/rows"]
)
@pytest.mark.asyncio
async def test_post_export_rejects_anonymous(client, path):
    resp = await client.post(path, json={"rows": [_api_row()], "to_email": "a@b.ch"})
    assert resp.status_code in (401, 403)
    assert "error" in resp.json()


@pytest.mark.asyncio
async def test_post_banana_returns_utf8_tsv_attachment(client, headers):
    resp = await client.post(
        "/api/export/banana", json={"rows": [_api_row(), _api_row(nr=2, betrag=-3.333)]}, headers=headers
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/plain")
    assert resp.headers["content-disposition"] == "attachment; filename=banana_import.txt"
    text = resp.content.decode("utf-8")
    lines = text.split("\n")
    assert lines[0] == BANANA_HEADER
    assert lines[1] == "2025-03-15\tCafé Müller\t6500\t1020\t12.50\tI81"
    assert lines[2].split("\t")[4] == "-3.33"


@pytest.mark.asyncio
async def test_post_csv_and_excel_return_attachments(client, headers):
    resp = await client.post("/api/export/csv", json={"rows": [_api_row()]}, headers=headers)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert "Café Müller" in resp.content.decode("utf-8")

    resp = await client.post("/api/export/excel", json={"rows": [_api_row()]}, headers=headers)
    assert resp.status_code == 200
    assert resp.headers["content-disposition"] == "attachment; filename=buchhaltung.xlsx"
    assert load_workbook(io.BytesIO(resp.content)).active.cell(row=2, column=5).value == "Café Müller"


@pytest.mark.parametrize("fmt", ["banana", "csv", "excel"])
@pytest.mark.asyncio
async def test_post_export_with_no_rows_is_404(client, headers, fmt):
    resp = await client.post(f"/api/export/{fmt}", json={"rows": []}, headers=headers)
    assert resp.status_code == 404
    assert resp.json()["error"]["code"] == "http_404"


@pytest.mark.asyncio
async def test_post_export_mwstchf_accepts_blank_string(client, headers):
    resp = await client.post("/api/export/csv", json={"rows": [_api_row(mwstchf="")]}, headers=headers)
    assert resp.status_code == 200
    assert resp.content.decode("utf-8").strip().split("\n")[1].split(";")[11] == "0"


@pytest.mark.asyncio
async def test_get_banana_exports_db_bookings_in_id_order(client, db_session):
    tenant = await create_tenant(db_session)
    user = await create_user(db_session, tenant)
    await create_booking(db_session, tenant, datum="03.01.2025", beschreibung="Zürich Taxi", betrag=45.0, source="pdf")
    await create_booking(
        db_session, tenant, datum="04.01.2025", beschreibung="Gutschrift", betrag=-100.0, mwst_code="V81", source="pdf"
    )
    await create_booking(db_session, tenant, datum="05.01.2025", beschreibung="Scan", betrag=7.777, source="scanner")

    resp = await client.get("/api/export/banana", headers=auth_headers(user))
    assert resp.status_code == 200
    lines = resp.content.decode("utf-8").split("\n")
    assert lines[0] == BANANA_HEADER
    assert lines[1] == "2025-01-03\tZürich Taxi\t6500\t1020\t45.00\t"
    assert lines[2] == "2025-01-04\tGutschrift\t6500\t1020\t-100.00\tV81"
    assert lines[3] == "2025-01-05\tScan\t6500\t1020\t7.78\t"

    resp = await client.get("/api/export/banana", params={"source": "scanner"}, headers=auth_headers(user))
    assert resp.content.decode("utf-8").split("\n")[1:] == ["2025-01-05\tScan\t6500\t1020\t7.78\t"]


@pytest.mark.asyncio
async def test_get_export_with_no_bookings_is_404(client, db_session):
    tenant = await create_tenant(db_session)
    user = await create_user(db_session, tenant)

    for fmt in ("banana", "csv", "excel"):
        resp = await client.get(f"/api/export/{fmt}", headers=auth_headers(user))
        assert resp.status_code == 404, fmt
